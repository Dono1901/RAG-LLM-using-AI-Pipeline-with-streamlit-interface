"""
Local RAG Application using Ollama + Sentence Transformers.
No API keys required - runs entirely on your machine.
Enhanced with Excel processing and financial analysis capabilities.
"""

import hashlib
import logging
import os
import threading
from pathlib import Path
from typing import Any, Dict, List, Optional

import numpy as np

from dotenv import load_dotenv

from config import settings
from local_llm import LocalLLM, LocalEmbedder
from logging_config import setup_logging
from protocols import LLMProvider, EmbeddingProvider

# Setup structured logging
setup_logging()

logger = logging.getLogger(__name__)

load_dotenv()


class SimpleRAG:
    """Simple RAG implementation for local use with Excel and financial analysis support."""

    # Supported Excel extensions
    EXCEL_EXTENSIONS = {'.xlsx', '.xlsm', '.xls', '.csv', '.tsv'}

    def __init__(
        self,
        docs_folder: str = "./documents",
        llm_model: str = "llama3.2",
        embedding_model: str = "all-MiniLM-L6-v2",
        llm: Optional[LLMProvider] = None,
        embedder: Optional[EmbeddingProvider] = None,
        store=None,
    ):
        """
        Initialize the RAG system.

        Args:
            docs_folder: Path to folder containing PDF/text/Excel documents
            llm_model: Ollama model name (ignored if *llm* is provided)
            embedding_model: Sentence transformer model name (ignored if *embedder* is provided)
            llm: Optional pre-built LLM provider (for dependency injection / testing)
            embedder: Optional pre-built embedding provider (for dependency injection / testing)
            store: Optional Neo4jStore instance (auto-connects if NEO4J_URI is set and store is None)
        """
        self.docs_folder = Path(docs_folder)
        self.docs_folder.mkdir(exist_ok=True)

        if llm is not None:
            self.llm = llm
        else:
            logger.info(f"Initializing LocalLLM with model: {llm_model}")
            self.llm = LocalLLM(
                model=llm_model,
                timeout_seconds=settings.llm_timeout_seconds,
                max_retries=settings.llm_max_retries,
                circuit_breaker_failure_threshold=settings.circuit_breaker_failure_threshold,
                circuit_breaker_recovery_seconds=settings.circuit_breaker_recovery_seconds,
            )

        if embedder is not None:
            self.embedder = embedder
        else:
            logger.info(f"Initializing LocalEmbedder with model: {embedding_model}")
            self.embedder = LocalEmbedder(model_name=embedding_model)

        # Document store (simple in-memory)
        self.documents: List[Dict[str, Any]] = []
        self.embeddings: List[list] = []

        # Pre-computed numpy matrix and norms (built after document load)
        self._doc_matrix: Optional[np.ndarray] = None
        self._doc_norms: Optional[np.ndarray] = None

        # BM25 index (optional - if rank_bm25 is installed)
        self._bm25_index = None
        self._bm25_available = False
        try:
            from rank_bm25 import BM25Okapi
            self._bm25_available = True
            logger.info("BM25 search enabled")
        except ImportError:
            logger.info("BM25 not available (rank_bm25 not installed) - using semantic-only search")

        # Thread safety for reload
        self._lock = threading.Lock()

        # Store embedding model name for cache key stability
        self._embedding_model_name = embedding_model

        # Graph store (optional - activates when NEO4J_URI is set)
        if store is not None:
            self._graph_store = store
        else:
            try:
                from graph_store import Neo4jStore
                self._graph_store = Neo4jStore.connect()
            except ImportError:
                self._graph_store = None

        # Excel processor and financial analyzer (lazy loaded)
        self._excel_processor = None
        self._charlie_analyzer = None
        self._financial_analysis_cache = None

        # Per-period FinancialData cache for multi-document comparison
        self._period_financial_data: Dict[str, Any] = {}

        # Embedding cache directory
        self._cache_dir = Path(settings.embedding_cache_dir)
        self._cache_dir.mkdir(parents=True, exist_ok=True)

        # Load documents on init
        self._load_documents()

    @property
    def excel_processor(self):
        """Lazy-load Excel processor."""
        if self._excel_processor is None:
            try:
                from excel_processor import ExcelProcessor
                self._excel_processor = ExcelProcessor(str(self.docs_folder))
            except ImportError:
                logger.warning("Excel processor not available. Install openpyxl and pandas.")
        return self._excel_processor

    @property
    def charlie_analyzer(self):
        """Lazy-load financial analyzer."""
        if self._charlie_analyzer is None:
            try:
                from financial_analyzer import CharlieAnalyzer
                self._charlie_analyzer = CharlieAnalyzer()
            except ImportError:
                logger.warning("Financial analyzer not available.")
        return self._charlie_analyzer

    # ------------------------------------------------------------------
    # Financial analysis context for RAG
    # ------------------------------------------------------------------

    def _get_financial_analysis_context(self) -> str:
        """Compute and cache financial analysis from Excel data for LLM context.

        Scans loaded Excel files, runs CharlieAnalyzer on DataFrames,
        and formats results as concise text for prompt injection.

        Returns:
            Formatted analysis text, or empty string if unavailable.
        """
        if self._financial_analysis_cache is not None:
            return self._financial_analysis_cache

        if not self.charlie_analyzer or not self.excel_processor:
            self._financial_analysis_cache = ""
            return ""

        try:
            excel_files = self.excel_processor.scan_for_excel_files()
            if not excel_files:
                self._financial_analysis_cache = ""
                return ""

            analysis_parts = []

            for file_path in excel_files[:3]:
                try:
                    workbook = self.excel_processor.load_workbook(file_path)
                    combined = self.excel_processor.combine_sheets_intelligently(workbook)

                    # Try merged DataFrame first, then individual sheets
                    dfs_to_analyze = []
                    if combined.merged_df is not None and not combined.merged_df.empty:
                        dfs_to_analyze.append((file_path.name, combined.merged_df))
                    else:
                        for sheet in workbook.sheets[:3]:
                            if not sheet.df.empty:
                                dfs_to_analyze.append(
                                    (f"{file_path.name}/{sheet.name}", sheet.df)
                                )

                    for source_name, df in dfs_to_analyze:
                        financial_data = self.charlie_analyzer._dataframe_to_financial_data(df)

                        # Only analyze if we have meaningful data
                        if financial_data.revenue is None and financial_data.total_assets is None:
                            continue

                        # Cache for multi-document comparison
                        if hasattr(self, "_period_financial_data"):
                            self._period_financial_data[source_name] = financial_data

                        report = self.charlie_analyzer.generate_report(financial_data)

                        # Persist to graph if available (structured path preferred)
                        if getattr(self, "_graph_store", None):
                            try:
                                from graph_retriever import persist_structured_analysis_to_graph
                                from ratio_framework import run_all_ratios
                                ratio_results = run_all_ratios(financial_data)
                                persist_structured_analysis_to_graph(
                                    self._graph_store,
                                    source_name,
                                    source_name,
                                    financial_data=financial_data,
                                    ratio_results=ratio_results,
                                )
                            except Exception as exc:
                                logger.debug("Graph persist failed: %s", exc)

                        analysis_parts.append(
                            f"=== Computed Analysis: {source_name} ===\n"
                            f"{report.executive_summary}\n\n"
                            f"Key Ratios:\n{report.sections.get('ratio_analysis', 'N/A')}\n\n"
                            f"Scoring Models:\n{report.sections.get('scoring_models', 'N/A')}\n\n"
                            f"Risk Assessment:\n{report.sections.get('risk_assessment', 'N/A')}"
                        )
                        break  # One analysis per file

                except Exception as e:
                    logger.debug(f"Could not analyze {file_path.name}: {e}")

            self._financial_analysis_cache = "\n\n".join(analysis_parts) if analysis_parts else ""

        except Exception as e:
            logger.debug(f"Financial analysis context generation failed: {e}")
            self._financial_analysis_cache = ""

        return self._financial_analysis_cache

    # ------------------------------------------------------------------
    # Embedding cache helpers
    # ------------------------------------------------------------------

    def _embedding_cache_key(self, file_path: Path) -> str:
        """Generate a cache key from content hash + embedding model.

        Uses SHA-256 of file content instead of mtime, so unchanged files
        won't be re-embedded even after a touch/copy operation.
        """
        content_hash = hashlib.sha256(file_path.read_bytes()).hexdigest()
        key_str = f"{file_path.name}:{content_hash}:{self._embedding_model_name}"
        return hashlib.sha256(key_str.encode()).hexdigest()

    def _load_cached_embeddings(self, cache_key: str):
        """Load cached embeddings from disk via joblib, or return None."""
        cache_file = self._cache_dir / f"{cache_key}.joblib"
        if cache_file.exists():
            try:
                import joblib
                data = joblib.load(cache_file)
                logger.info(f"Loaded cached embeddings: {cache_key[:12]}...")
                return data  # (documents_list, embeddings_list)
            except Exception as e:
                logger.debug(f"Cache read failed: {e}")
        return None

    def _save_cached_embeddings(self, cache_key: str, documents: list, embeddings: list):
        """Save embeddings to disk cache."""
        cache_file = self._cache_dir / f"{cache_key}.joblib"
        try:
            import joblib
            joblib.dump((documents, embeddings), cache_file)
            logger.debug(f"Saved embeddings cache: {cache_key[:12]}...")
        except Exception as e:
            logger.debug(f"Cache write failed: {e}")

    # ------------------------------------------------------------------
    # Document loading
    # ------------------------------------------------------------------

    def _extract_text(self, file_path: Path) -> Optional[str]:
        """Extract text from a PDF or text file. Returns None on failure."""
        suffix = file_path.suffix.lower()

        if suffix == ".pdf":
            import fitz
            doc = fitz.open(file_path)
            text = "".join(page.get_text() for page in doc)
            doc.close()
            return text if text.strip() else None

        if suffix in (".txt", ".md"):
            text = file_path.read_text(encoding="utf-8")
            return text if text.strip() else None

        if suffix == ".docx":
            from docx import Document as DocxDocument
            doc = DocxDocument(file_path)
            text = "\n".join(p.text for p in doc.paragraphs)
            return text if text.strip() else None

        return None

    def _load_documents(self):
        """Load and process documents from the docs folder (PDF, TXT, MD, Excel, CSV)."""
        logger.info(f"Loading documents from {self.docs_folder}")

        for file_path in self.docs_folder.rglob("*"):
            if not file_path.is_file():
                continue
            file_size = file_path.stat().st_size
            if file_size > settings.max_file_size_mb * 1024 * 1024:
                logger.warning(
                    f"Skipping {file_path} ({file_size / 1024 / 1024:.1f} MB) "
                    f"- exceeds {settings.max_file_size_mb} MB limit"
                )
                continue
            try:
                suffix = file_path.suffix.lower()

                # Try embedding cache for this file
                cache_key = self._embedding_cache_key(file_path)
                cached = self._load_cached_embeddings(cache_key)
                if cached is not None:
                    docs, embs = cached
                    self.documents.extend(docs)
                    self.embeddings.extend(embs)
                    continue

                file_docs: List[Dict[str, Any]] = []

                if suffix in (".pdf", ".txt", ".md", ".docx"):
                    text = self._extract_text(file_path)
                    if text:
                        if suffix == ".docx":
                            doc_type = "docx"
                        elif suffix == ".pdf":
                            doc_type = "pdf"
                        else:
                            doc_type = "text"
                        chunks = self._chunk_text(
                            text,
                            chunk_size=settings.chunk_size,
                            overlap=settings.chunk_overlap,
                        )
                        rel_path = str(file_path.relative_to(self.docs_folder))
                        for chunk in chunks:
                            file_docs.append({
                                "source": rel_path,
                                "content": chunk,
                                "type": doc_type,
                            })
                        logger.info(f"Loaded {len(chunks)} chunks from {rel_path}")

                elif suffix in self.EXCEL_EXTENSIONS:
                    excel_docs = self._process_excel_file(file_path)
                    file_docs.extend(excel_docs)
                    logger.info(f"Loaded {len(excel_docs)} chunks from Excel file {file_path.name}")

                # Generate and cache embeddings for this file's docs
                if file_docs:
                    texts = [d["content"] for d in file_docs]
                    embs = self.embedder.embed_batch(texts)
                    self._save_cached_embeddings(cache_key, file_docs, embs)
                    self.documents.extend(file_docs)
                    self.embeddings.extend(embs)

                    # Persist to Neo4j graph store (if available)
                    if getattr(self, "_graph_store", None):
                        self._graph_store.store_chunks(
                            chunks=file_docs,
                            embeddings=embs,
                            doc_id=file_docs[0].get("source", str(file_path.name)),
                        )

            except Exception as e:
                logger.warning(f"Failed to load {file_path}: {e}")

        if self.documents:
            logger.info(f"Total: {len(self.documents)} chunks indexed")
            self._build_embedding_index()
            self._build_bm25_index()

            # Ensure Neo4j schema (vector index) after first load
            if getattr(self, "_graph_store", None) and self.embeddings:
                try:
                    dim = len(self.embeddings[0])
                    self._graph_store.ensure_schema(dim, self._embedding_model_name)
                except Exception as exc:
                    logger.warning("Neo4j schema setup failed: %s", exc)
        else:
            logger.warning("No documents found in the documents folder")

    def _build_embedding_index(self):
        """Pre-compute the numpy doc matrix and norms for fast retrieval."""
        if not self.embeddings:
            self._doc_matrix = None
            self._doc_norms = None
            return
        self._doc_matrix = np.asarray(self.embeddings, dtype=np.float32)
        norms = np.linalg.norm(self._doc_matrix, axis=1)
        # Replace zero norms with 1.0 to avoid division by zero
        self._doc_norms = np.where(norms == 0, 1.0, norms)

    def _build_bm25_index(self):
        """Build BM25 index for keyword search."""
        if not self._bm25_available or not self.documents:
            self._bm25_index = None
            return

        try:
            from rank_bm25 import BM25Okapi
            # Tokenize all document contents (simple whitespace + lowercase)
            tokenized_corpus = [
                doc["content"].lower().split() for doc in self.documents
            ]
            self._bm25_index = BM25Okapi(tokenized_corpus)
            logger.info("BM25 index built successfully")
        except Exception as e:
            logger.warning(f"Failed to build BM25 index: {e}")
            self._bm25_index = None

    def _process_excel_file(self, file_path: Path) -> List[Dict[str, Any]]:
        """Process Excel file into RAG-compatible document chunks."""
        if self.excel_processor is None:
            logger.warning(f"Excel processor not available, skipping {file_path.name}")
            return []

        try:
            from excel_processor import process_excel_for_rag
            return process_excel_for_rag(file_path, str(self.docs_folder))
        except Exception as e:
            logger.error(f"Failed to process Excel file {file_path}: {e}")
            return []

    def _chunk_text(self, text: str, chunk_size: int = 500, overlap: int = 50) -> list:
        """Split text into overlapping chunks."""
        words = text.split()
        chunks = []

        for i in range(0, len(words), chunk_size - overlap):
            chunk = " ".join(words[i:i + chunk_size])
            if chunk.strip():
                chunks.append(chunk)

        return chunks if chunks else [text]

    def _cosine_similarity(self, vec1: list, vec2: list) -> float:
        """Calculate cosine similarity between two vectors."""
        a = np.asarray(vec1, dtype=np.float32)
        b = np.asarray(vec2, dtype=np.float32)
        norm_a = np.linalg.norm(a)
        norm_b = np.linalg.norm(b)
        if norm_a == 0 or norm_b == 0:
            return 0.0
        return float(np.dot(a, b) / (norm_a * norm_b))

    def retrieve(self, query: str, top_k: int = 3) -> list:
        """
        Retrieve the most relevant documents for a query using hybrid search.

        Combines semantic (embedding-based) and keyword (BM25) search using
        Reciprocal Rank Fusion (RRF) if BM25 is available, otherwise falls
        back to semantic-only search.

        Args:
            query: User's question
            top_k: Number of documents to retrieve

        Returns:
            List of relevant document chunks
        """
        if not self.documents:
            return []

        # Clamp top_k to valid range
        top_k = max(1, min(top_k, settings.max_top_k, len(self.documents)))

        # If BM25 is not available, fall back to semantic-only search
        if not self._bm25_available or self._bm25_index is None:
            return self._semantic_search(query, top_k)

        # Perform both semantic and BM25 search
        semantic_results = self._semantic_search(query, top_k * 2)  # Get more for fusion
        bm25_results = self._bm25_search(query, top_k * 2)

        # Fuse results using Reciprocal Rank Fusion
        fused_results = self._fuse_results_rrf(semantic_results, bm25_results, top_k)
        return fused_results

    def _semantic_search(self, query: str, top_k: int) -> list:
        """
        Perform semantic search using embeddings.

        Tries Neo4j vector index first (if available), falls back to in-memory numpy.

        Args:
            query: User's question
            top_k: Number of documents to retrieve

        Returns:
            List of relevant document chunks
        """
        if not self.documents:
            return []

        # Clamp top_k to valid range
        top_k = max(1, min(top_k, len(self.documents)))

        # Embed the query
        query_embedding = self.embedder.embed(query)

        # Try Neo4j graph search first (vector + graph traversal for enriched context)
        if getattr(self, "_graph_store", None):
            try:
                neo4j_results = self._graph_store.graph_search(
                    query_embedding, top_k, self._embedding_model_name,
                )
                if neo4j_results:
                    return [
                        {
                            "source": r.get("source", ""),
                            "content": r.get("content", ""),
                            "type": "unknown",
                            "_graph_context": {
                                "document": r.get("document", ""),
                                "period": r.get("period", ""),
                                "ratios": r.get("ratios", []),
                                "scores": r.get("scores", []),
                            },
                        }
                        for r in neo4j_results
                    ]
            except Exception as exc:
                logger.debug("Neo4j graph search failed, falling back to numpy: %s", exc)

        # In-memory numpy fallback
        query_vec = np.asarray(query_embedding, dtype=np.float32)
        query_norm = np.linalg.norm(query_vec)

        if query_norm == 0:
            return self.documents[:top_k]

        # Use pre-computed matrix and norms (falls back to on-the-fly if not built)
        if self._doc_matrix is not None and self._doc_norms is not None:
            similarities = self._doc_matrix @ query_vec / (self._doc_norms * query_norm)
        else:
            doc_matrix = np.asarray(self.embeddings, dtype=np.float32)
            norms = np.linalg.norm(doc_matrix, axis=1)
            safe_norms = np.where(norms == 0, 1.0, norms)
            similarities = doc_matrix @ query_vec / (safe_norms * query_norm)

        # Use argpartition for O(n) partial sort when doc count >> top_k
        if len(similarities) > top_k * 4:
            partitioned = np.argpartition(similarities, -top_k)[-top_k:]
            top_indices = partitioned[np.argsort(similarities[partitioned])][::-1]
        else:
            top_indices = np.argsort(similarities)[-top_k:][::-1]

        return [self.documents[i] for i in top_indices]

    def _bm25_search(self, query: str, top_k: int) -> list:
        """
        Perform BM25 keyword search.

        Args:
            query: User's question
            top_k: Number of documents to retrieve

        Returns:
            List of relevant document chunks
        """
        if not self._bm25_index or not self.documents:
            return []

        # Clamp top_k to valid range
        top_k = max(1, min(top_k, len(self.documents)))

        # Tokenize query (same way as corpus)
        tokenized_query = query.lower().split()

        # Get BM25 scores
        scores = self._bm25_index.get_scores(tokenized_query)

        # Get top-k indices
        if len(scores) > top_k * 4:
            partitioned = np.argpartition(scores, -top_k)[-top_k:]
            top_indices = partitioned[np.argsort(scores[partitioned])][::-1]
        else:
            top_indices = np.argsort(scores)[-top_k:][::-1]

        return [self.documents[i] for i in top_indices]

    def _fuse_results_rrf(
        self,
        semantic_results: list,
        bm25_results: list,
        top_k: int
    ) -> list:
        """
        Fuse semantic and BM25 results using Reciprocal Rank Fusion.

        RRF formula: score(d) = sum over systems s of: weight_s / (k + rank_s(d))
        where k is typically 60 (default RRF constant).

        Args:
            semantic_results: Results from semantic search
            bm25_results: Results from BM25 search
            top_k: Number of final results to return

        Returns:
            Fused list of document chunks
        """
        # Build rank maps (doc_id -> rank) for each system
        semantic_ranks = {id(doc): rank for rank, doc in enumerate(semantic_results)}
        bm25_ranks = {id(doc): rank for rank, doc in enumerate(bm25_results)}

        # Get all unique documents
        all_docs = {}
        for doc in semantic_results + bm25_results:
            all_docs[id(doc)] = doc

        # Calculate RRF scores
        rrf_scores = {}
        k = settings.rrf_k

        for doc_id, doc in all_docs.items():
            score = 0.0

            # Semantic contribution
            if doc_id in semantic_ranks:
                score += settings.semantic_weight / (k + semantic_ranks[doc_id])

            # BM25 contribution
            if doc_id in bm25_ranks:
                score += settings.bm25_weight / (k + bm25_ranks[doc_id])

            rrf_scores[doc_id] = score

        # Sort by RRF score and return top-k
        sorted_docs = sorted(
            all_docs.items(),
            key=lambda x: rrf_scores[x[0]],
            reverse=True
        )[:top_k]

        return [doc for doc_id, doc in sorted_docs]

    def answer(self, query: str, retrieved_docs: Optional[List[Dict[str, Any]]] = None) -> str:
        """
        Answer a question using RAG with financial analysis support.

        Args:
            query: User's question
            retrieved_docs: Pre-retrieved documents (avoids double retrieval)

        Returns:
            Generated answer
        """
        # Enforce query length limit
        if len(query) > settings.max_query_length:
            return f"Query too long. Maximum length is {settings.max_query_length} characters."

        # Retrieve relevant documents (reuse if already provided)
        relevant_docs = retrieved_docs if retrieved_docs is not None else self.retrieve(query, top_k=settings.top_k)

        if not relevant_docs:
            return "No documents loaded. Please add PDF, text, or Excel files to the 'documents' folder."

        # Check if query involves financial analysis
        is_financial = self._is_financial_query(query)

        # Build context from retrieved documents
        context_parts = []
        excel_data = []

        for doc in relevant_docs:
            source_info = f"[Source: {doc['source']}]"
            doc_type = doc.get('type', 'unknown')

            if doc_type == 'excel':
                financial_type = doc.get('metadata', {}).get('financial_type', '')
                if financial_type:
                    source_info += f" [Type: {financial_type}]"
                excel_data.append(doc)

            context_parts.append(f"{source_info}\n{doc['content']}")

        context = "\n\n".join(context_parts)

        # Enhanced prompt for financial queries (excel data or graph context)
        has_graph_context = any(d.get("_graph_context") for d in relevant_docs)
        if is_financial and self.charlie_analyzer and (excel_data or has_graph_context):
            prompt = self._build_financial_prompt(query, context, excel_data, relevant_docs)
        else:
            # Standard RAG prompt
            prompt = f"""You are a helpful assistant that answers questions based on the provided context.
Use ONLY the information from the context to answer. If the answer is not in the context, say so.

Context:
{context}

Question: {query}

Answer:"""

        # Generate answer
        try:
            answer = self.llm.generate(prompt)
            return answer.strip()
        except Exception as e:
            return f"Error generating answer: {e}"

    def retrieve_with_decomposition(self, query: str, top_k: int = 3) -> list:
        """Retrieve using query decomposition for improved recall.

        Expands the original query into multiple sub-queries targeting
        different aspects, retrieves for each, and deduplicates results
        using source+content identity.

        Args:
            query: User's question
            top_k: Number of final documents to return

        Returns:
            Deduplicated list of relevant document chunks
        """
        sub_queries = self._decompose_query(query)

        # Retrieve for each sub-query
        all_results = []
        seen_keys = set()

        for sq in sub_queries:
            results = self.retrieve(sq, top_k=top_k)
            for doc in results:
                # Deduplicate by source + first 200 chars of content
                key = (doc.get('source', ''), doc.get('content', '')[:200])
                if key not in seen_keys:
                    seen_keys.add(key)
                    all_results.append(doc)

        return all_results[:top_k]

    def _decompose_query(self, query: str) -> list:
        """Decompose a complex query into sub-queries for multi-query retrieval.

        Uses keyword-based expansion (no LLM call) to generate variations
        that capture different aspects of the question.

        Args:
            query: Original user query

        Returns:
            List of sub-queries including the original
        """
        sub_queries = [query]  # Always include original
        query_lower = query.lower()

        # Financial aspect expansions
        expansions = {
            'profitability': ['revenue', 'net income', 'margin', 'profit'],
            'liquidity': ['current ratio', 'cash', 'working capital'],
            'leverage': ['debt', 'equity', 'interest coverage'],
            'efficiency': ['asset turnover', 'inventory turnover', 'receivables'],
            'growth': ['revenue growth', 'trend', 'year over year'],
            'risk': ['z-score', 'bankruptcy', 'distress', 'leverage'],
            'cash flow': ['operating cash flow', 'free cash flow', 'capex'],
            'valuation': ['roe', 'roa', 'roic', 'earnings'],
        }

        # Find which aspects the query touches
        matched_aspects = []
        for aspect, keywords in expansions.items():
            if any(kw in query_lower for kw in keywords) or aspect in query_lower:
                matched_aspects.append(aspect)

        # Generate sub-queries for related aspects
        if matched_aspects:
            for aspect in matched_aspects:
                for keyword in expansions[aspect]:
                    if keyword not in query_lower:
                        sub_queries.append(f"{query} {keyword}")
                        if len(sub_queries) >= 4:  # Cap at 4 sub-queries
                            return sub_queries

        # If no financial aspects matched, add simple reformulations
        if len(sub_queries) == 1:
            # Add a specificity variant
            if '?' in query:
                sub_queries.append(query.replace('?', ' details?'))
            else:
                sub_queries.append(f"details about {query}")

        return sub_queries[:4]

    def _is_financial_query(self, query: str) -> bool:
        """Check if query involves financial analysis."""
        financial_keywords = [
            'ratio', 'margin', 'profit', 'revenue', 'income', 'expense',
            'cash flow', 'budget', 'variance', 'roe', 'roa', 'roi',
            'liquidity', 'leverage', 'debt', 'equity', 'asset', 'liability',
            'growth', 'trend', 'forecast', 'analysis', 'financial',
            'balance sheet', 'income statement', 'p&l', 'cfo', 'ebitda',
            'z-score', 'zscore', 'f-score', 'fscore', 'piotroski', 'altman',
            'dupont', 'health score', 'composite', 'working capital',
            'bankruptcy', 'distress', 'scoring', 'grade',
        ]
        query_lower = query.lower()
        return any(keyword in query_lower for keyword in financial_keywords)

    @staticmethod
    def _is_temporal_comparison_query(query: str) -> bool:
        """Detect queries that compare across time periods."""
        temporal_patterns = [
            "change from", "changed from", "year over year", "yoy",
            "quarter over quarter", "qoq", "compared to", "comparison",
            "trend", "trends", "over time", "growth rate",
            "improved", "deteriorated", "worsened", "increased", "decreased",
            "vs ", "versus", "relative to", "from fy", "from q",
        ]
        query_lower = query.lower()
        return any(pattern in query_lower for pattern in temporal_patterns)

    def _build_financial_prompt(
        self,
        query: str,
        context: str,
        excel_data: List[Dict],
        relevant_docs: Optional[List[Dict[str, Any]]] = None,
    ) -> str:
        """Build an enhanced prompt for financial queries with computed analysis."""
        # Get pre-computed analysis results (cached after first call)
        computed_analysis = self._get_financial_analysis_context()

        # Build the computed analysis section
        computed_section = ""
        if computed_analysis:
            computed_section = f"""
COMPUTED FINANCIAL ANALYSIS (use these exact values when answering):
{computed_analysis}
"""

        # Build graph-retrieved metrics section
        graph_section = ""
        if relevant_docs:
            graph_contexts = [
                d["_graph_context"]
                for d in relevant_docs
                if d.get("_graph_context")
                and (d["_graph_context"].get("ratios") or d["_graph_context"].get("scores"))
            ]
            if graph_contexts:
                from graph_retriever import format_graph_context
                formatted = format_graph_context(graph_contexts)
                if formatted:
                    graph_section = f"""
GRAPH-RETRIEVED FINANCIAL METRICS:
{formatted}
"""

        return f"""You are a senior financial analyst with CFO-level expertise, inspired by Charlie Munger's
analytical framework. Focus on fundamentals, cash flow, and sustainable competitive advantages.

USER QUERY: {query}

FINANCIAL DATA CONTEXT:
{context}
{graph_section}{computed_section}
ANALYSIS FRAMEWORK (Charlie Munger approach):
1. What are the key value drivers?
2. What could go wrong? (Inversion thinking)
3. Is there a margin of safety?
4. Focus on cash flows, not just accounting profits
5. Look for sustainable competitive advantages

Provide a comprehensive, actionable analysis addressing the query.
When computed ratios, scores, or grades are available above, cite those EXACT values.
Include specific numbers from the data and cite your sources.
If performing calculations, show your work.

Answer:"""

    def answer_stream(self, query: str, retrieved_docs: Optional[List[Dict[str, Any]]] = None):
        """Stream an answer using RAG with financial analysis support.

        Falls back to non-streaming if the LLM doesn't support streaming.

        Args:
            query: User's question
            retrieved_docs: Pre-retrieved documents (avoids double retrieval)

        Yields:
            str: Text chunks as they arrive from the LLM
        """
        if len(query) > settings.max_query_length:
            yield f"Query too long. Maximum length is {settings.max_query_length} characters."
            return

        relevant_docs = retrieved_docs if retrieved_docs is not None else self.retrieve(query, top_k=settings.top_k)

        if not relevant_docs:
            yield "No documents loaded. Please add PDF, text, or Excel files to the 'documents' folder."
            return

        is_financial = self._is_financial_query(query)

        context_parts = []
        excel_data = []

        for doc in relevant_docs:
            source_info = f"[Source: {doc['source']}]"
            doc_type = doc.get('type', 'unknown')

            if doc_type == 'excel':
                financial_type = doc.get('metadata', {}).get('financial_type', '')
                if financial_type:
                    source_info += f" [Type: {financial_type}]"
                excel_data.append(doc)

            context_parts.append(f"{source_info}\n{doc['content']}")

        context = "\n\n".join(context_parts)

        has_graph_context = any(d.get("_graph_context") for d in relevant_docs)
        if is_financial and self.charlie_analyzer and (excel_data or has_graph_context):
            prompt = self._build_financial_prompt(query, context, excel_data, relevant_docs)
        else:
            prompt = f"""You are a helpful assistant that answers questions based on the provided context.
Use ONLY the information from the context to answer. If the answer is not in the context, say so.

Context:
{context}

Question: {query}

Answer:"""

        try:
            if hasattr(self.llm, 'generate_stream'):
                yield from self.llm.generate_stream(prompt)
            else:
                answer = self.llm.generate(prompt)
                yield answer.strip()
        except Exception as e:
            yield f"Error generating answer: {e}"

    def reload_documents(self):
        """Reload documents from the folder (thread-safe)."""
        with self._lock:
            self.documents = []
            self.embeddings = []
            self._doc_matrix = None
            self._doc_norms = None
            self._bm25_index = None
            self._financial_analysis_cache = None
            self._period_financial_data = {}
            self._load_documents()


# CLI-only entry point (not used by Streamlit - see streamlit_app_local.py)
if __name__ == "__main__":
    rag = SimpleRAG(
        docs_folder="./documents",
        llm_model=os.getenv("OLLAMA_MODEL", settings.llm_model),
        embedding_model=os.getenv("EMBEDDING_MODEL", settings.embedding_model),
    )

    print("\n" + "=" * 50)
    print("Local RAG System Ready!")
    print("=" * 50)
    print(f"Documents loaded: {len(rag.documents)} chunks")
    print(f"Documents folder: {rag.docs_folder.absolute()}")
    print("\nAdd PDF or TXT files to the 'documents' folder to get started.")
    print("=" * 50 + "\n")

    while True:
        query = input("\nAsk a question (or 'quit' to exit): ").strip()
        if query.lower() in ["quit", "exit", "q"]:
            break
        if query:
            print("\nSearching and generating answer...")
            answer = rag.answer(query)
            print(f"\nAnswer: {answer}")
